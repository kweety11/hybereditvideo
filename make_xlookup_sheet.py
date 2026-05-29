import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "XLOOKUP"
ws.sheet_view.rightToLeft = True  # عرض الورقة من اليمين لليسار

# ---- ألوان وتنسيقات ----
header_fill = PatternFill("solid", fgColor="2F5597")
header_font = Font(color="FFFFFF", bold=True, size=12)
box_fill = PatternFill("solid", fgColor="FFF2CC")
result_fill = PatternFill("solid", fgColor="E2EFDA")
title_font = Font(bold=True, size=14, color="2F5597")
center = Alignment(horizontal="center", vertical="center")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# ---- جدول البيانات (المصدر) ----
data = [
    ["كود المنتج", "اسم المنتج", "السعر", "الكمية", "القسم"],
    [1001, "لابتوب", 18500, 12, "إلكترونيات"],
    [1002, "ماوس", 250,  80, "إكسسوارات"],
    [1003, "كيبورد", 600, 45, "إكسسوارات"],
    [1004, "شاشة", 4200, 20, "إلكترونيات"],
    [1005, "سماعة", 1500, 35, "صوتيات"],
    [1006, "طابعة", 3100, 15, "إلكترونيات"],
    [1007, "كاميرا", 9800, 8,  "تصوير"],
]

ws["A1"] = "جدول المنتجات (البيانات)"
ws["A1"].font = title_font

start_row = 2
for r, row in enumerate(data, start=start_row):
    for c, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.border = border
        cell.alignment = center
        if r == start_row:
            cell.fill = header_fill
            cell.font = header_font

# جدول البيانات من A2:E9 ، صف العناوين في الصف 2 ، البيانات A3:E9
# ---- قسم البحث بـ XLOOKUP ----
ws["G1"] = "ابحث هنا باستخدام XLOOKUP"
ws["G1"].font = title_font

labels = [
    ("G3", "اكتب كود المنتج:"),
    ("G4", "اسم المنتج:"),
    ("G5", "السعر:"),
    ("G6", "الكمية:"),
    ("G7", "القسم:"),
]
for ref, text in labels:
    ws[ref] = text
    ws[ref].font = Font(bold=True)

# خلية الإدخال
ws["H3"] = 1004
ws["H3"].fill = box_fill
ws["H3"].alignment = center
ws["H3"].font = Font(bold=True, size=12)
ws["H3"].border = border

# صيغ XLOOKUP — تبحث عن الكود في H3 داخل عمود الأكواد A3:A9 وترجّع العمود المطلوب
ws["H4"] = '=XLOOKUP(H3, A3:A9, B3:B9, "غير موجود")'   # الاسم
ws["H5"] = '=XLOOKUP(H3, A3:A9, C3:C9, "غير موجود")'   # السعر
ws["H6"] = '=XLOOKUP(H3, A3:A9, D3:D9, "غير موجود")'   # الكمية
ws["H7"] = '=XLOOKUP(H3, A3:A9, E3:E9, "غير موجود")'   # القسم

for ref in ["H4", "H5", "H6", "H7"]:
    ws[ref].fill = result_fill
    ws[ref].alignment = center
    ws[ref].border = border

# قائمة منسدلة لاختيار الكود بسهولة
dv = DataValidation(type="list", formula1="=$A$3:$A$9", allow_blank=True)
ws.add_data_validation(dv)
dv.add(ws["H3"])

# ---- شرح سريع ----
ws["G10"] = "الشرح:"
ws["G10"].font = Font(bold=True, color="2F5597")
ws["G11"] = "غيّر الكود في الخلية H3 (أو اختاره من القائمة) وكل البيانات تتحدّث تلقائيًا."
ws["G12"] = "صيغة الاسم مثلاً:  =XLOOKUP(H3, A3:A9, B3:B9, \"غير موجود\")"
ws["G13"] = "H3 = القيمة المطلوبة | A3:A9 = عمود البحث | B3:B9 = عمود النتيجة"

# ---- عرض الأعمدة ----
widths = {"A": 12, "B": 14, "C": 10, "D": 9, "E": 14, "F": 3,
          "G": 18, "H": 16}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

out = "/home/user/hybereditvideo/XLOOKUP_Sheet.xlsx"
wb.save(out)
print("saved:", out)
